"""Build the 20-minute ESML / MMPROPS deck (PPTX) and PNG previews.

The preview renderer mirrors the PPTX elements closely enough for layout QA:
same 16:9 canvas, same inches/points, same colors and fonts.
"""

from __future__ import annotations

import pathlib
from dataclasses import dataclass, field

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HERE = pathlib.Path(__file__).resolve().parent
OUT_PPTX = HERE / "ESML-MMPROPS-20min.pptx"
PREVIEW = HERE / "preview"

# ── palette ────────────────────────────────────────────────────────────────
DARK = "0B3B2E"
DARK2 = "10251C"
ACCENT = "1F8A5B"
ACCENT2 = "2FAE73"
MINT = "EAF6F0"
INK = "17211C"
GRAY = "5C6B63"
LIGHT = "F4F7F5"
WHITE = "FFFFFF"
MINT_TEXT = "A9D6C0"

CJK = "Noto Sans CJK SC"
MONO = "DejaVu Sans Mono"

CJK_REG = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
CJK_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
MONO_TTF = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"

SCALE = 120  # px per inch → 1600 × 900


# ── element helpers ────────────────────────────────────────────────────────
def text(x, y, w, h, s, size=16, bold=False, color=INK, align="l", font="cjk", spacing=1.15):
    return {"t": "text", "x": x, "y": y, "w": w, "h": h, "text": s, "size": size,
            "bold": bold, "color": color, "align": align, "font": font, "spacing": spacing}


def rect(x, y, w, h, fill=MINT, radius=0.10, line=None, line_w=1.0):
    return {"t": "rect", "x": x, "y": y, "w": w, "h": h, "fill": fill,
            "radius": radius, "line": line, "line_w": line_w}


def code(x, y, w, h, lines, size=10.5):
    return {"t": "code", "x": x, "y": y, "w": w, "h": h, "lines": lines, "size": size}


def image(x, y, w, h, path):
    return {"t": "image", "x": x, "y": y, "w": w, "h": h, "path": str(path)}


def table(x, y, w, h, rows, widths, size=12, header=True):
    return {"t": "table", "x": x, "y": y, "w": w, "h": h, "rows": rows,
            "widths": widths, "size": size, "header": header}


def line(x1, y1, x2, y2, color=ACCENT, width=1.25):
    return {"t": "line", "x1": x1, "y1": y1, "x2": x2, "y2": y2,
            "color": color, "width": width}


def footer(n, dark=False):
    c = "7FA99A" if dark else "8A9991"
    return [
        line(0.75, 6.95, 12.58, 6.95, color="2C5A49" if dark else "DCE6E0", width=0.75),
        text(0.75, 7.05, 8, 0.25, "Ai4Energy · ESML / MMPROPS", 9, color=c),
        text(11.6, 7.05, 1.0, 0.25, f"{n:02d}", 9, color=c, align="r", font="mono"),
    ]


def head(label, title):
    return [
        text(0.75, 0.35, 8, 0.28, label, 11, bold=True, color=ACCENT),
        text(0.75, 0.68, 11.8, 0.75, title, 28, bold=True, color=DARK, spacing=1.05),
        line(0.75, 1.55, 1.55, 1.55, color=ACCENT2, width=2.0),
    ]


def slide(bg, elements, notes=""):
    return {"bg": bg, "elements": elements, "notes": notes}


# ── deck ───────────────────────────────────────────────────────────────────
def build_slides():
    S = []

    S.append(slide(DARK, [
        rect(0.9, 1.35, 0.9, 0.06, fill=ACCENT2, radius=0.0),
        text(0.9, 1.0, 8, 0.28, "Ai4Energy · 20 分钟技术分享", 12, color=MINT_TEXT),
        text(0.9, 1.55, 11.5, 1.2, "ESML × MMPROPS", 46, bold=True, color=WHITE),
        text(0.9, 2.9, 11.0, 0.8, "从分子到电网：能源系统的建模与物性计算", 24, color="D7EFE2"),
        line(0.9, 4.05, 12.4, 4.05, color="2C5A49", width=0.75),
        text(0.9, 4.3, 11.4, 0.6, "声明式模型 · 编译求解 · 独立物性引擎 · MIT 客户端", 16, color=MINT_TEXT),
        text(0.9, 6.6, 5, 0.25, "2026 · Ai4Energy", 11, color="7FA99A"),
    ], "开场 40 秒。先给结论：今天讲两个引擎，一个把模型变成可解系统，一个把物性变成可调用服务。不要念标题，直接说“从分子到电网”这句话。"))

    S.append(slide(WHITE, [
        *head("01 / 为什么", "两道坎：模型怎么变成可解系统，物性从哪里来"),
        rect(0.75, 2.0, 5.7, 2.9),
        text(1.05, 2.25, 5.1, 0.4, "建模", 20, bold=True, color=ACCENT),
        text(1.05, 2.8, 5.1, 1.2, "手写 Newton、稀疏雅可比、初值、事件……\n每换一个模型，重来一遍。", 15, color=INK, spacing=1.3),
        text(1.05, 4.05, 5.1, 0.3, "结果：换模型不再重写求解器。", 12, bold=True, color=ACCENT),
        rect(6.9, 2.0, 5.7, 2.9),
        text(7.2, 2.25, 5.1, 0.4, "物性", 20, bold=True, color=ACCENT),
        text(7.2, 2.8, 5.1, 1.2, "状态点、饱和、闪蒸、相平衡；\n关联式、单位、收敛各一套。", 15, color=INK, spacing=1.3),
        text(7.2, 4.05, 5.1, 0.3, "结果：一次调用拿到带单位的值。", 12, bold=True, color=ACCENT),
        text(0.75, 5.35, 11.8, 0.5, "我们做了两个引擎：一个管模型，一个管物性。", 20, bold=True, color=DARK),
        *footer(2),
    ], "2 分钟。用听众自己的痛点开场：建模和物性各自有一套坑。强调不是“又一个求解器”，而是把重复劳动收进编译链。"))

    S.append(slide(WHITE, [
        *head("02 / ESML", "把模型写成方程，剩下的交给编译器"),
        *[
            rect(0.7 + i * 2.45, 2.2, 2.15, 1.55, fill=MINT if i % 2 == 0 else LIGHT)
            for i in range(5)
        ],
        *[
            text(0.82 + i * 2.45, 2.38, 1.9, 0.35, t, 14, bold=True, color=DARK)
            for i, t in enumerate(["声明式模型", "展平 · 约简", "结构分析", "因果化 · 撕裂", "求解 · 积分"])
        ],
        *[
            text(0.82 + i * 2.45, 2.85, 1.9, 0.6, t, 11, color=GRAY)
            for i, t in enumerate(["XML / 方程", "DAE / 指标", "BLT / 分块", "显式步 / 隐式块", "牛顿 / BDF"])
        ],
        *[text(2.62 + i * 2.45, 2.72, 0.4, 0.4, "→", 18, color=ACCENT2, align="c") for i in range(4)],
        text(0.75, 4.25, 11.8, 0.45, "同一份模型，稳态与动态共用一条编译链。", 17, bold=True, color=DARK),
        text(0.75, 4.85, 11.8, 0.6, "结构、单位、因果在编译期检查；不是运行时猜。", 14, color=GRAY),
        rect(0.75, 5.55, 11.8, 0.85, fill=LIGHT),
        text(1.0, 5.72, 5.4, 0.3, "编译期：结构 / 单位 / 因果", 12.5, bold=True, color=ACCENT),
        text(6.5, 5.72, 5.8, 0.3, "运行时：数值 / 事件 / 积分", 12.5, bold=True, color=DARK),
        *footer(3),
    ], "1 分 40 秒。从左到右讲五步，每步一句话；重点落在“编译链”而不是单个求解器。"))

    S.append(slide(WHITE, [
        *head("02 / ESML", "一个模型长什么样"),
        code(0.75, 1.75, 7.1, 3.6, [
            '<modelType name="Decay">',
            '  <parameters>',
            '    <parameter name="lam" type="real"',
            '               default="0.5"/>',
            '  </parameters>',
            '  <variables>',
            '    <variable name="x" type="real"',
            '              role="state" initial="1.0"/>',
            '  </variables>',
            '  <equations>',
            '    <equation>Der(x) = -lam * x</equation>',
            '  </equations>',
            '</modelType>',
        ], 10.5),
        text(8.15, 1.95, 4.3, 0.4, "纯声明", 16, bold=True, color=ACCENT),
        text(8.15, 2.45, 4.3, 2.8, "• 变量 / 参数 / 方程\n• 组件、连接、数组与序集\n• 编译期做结构检查\n• 同一模型可稳态、可动态", 14, color=INK, spacing=1.45),
        text(0.75, 5.55, 11.8, 0.3, "模型是数据；求解器是编译器的一部分。", 12, color=GRAY),
        *footer(4),
    ], "1 分 20 秒。不要逐行念 XML；指出三件事：声明式、组件化、编译期检查。让听众知道模型不是脚本。"))

    S.append(slide(WHITE, [
        *head("02 / ESML", "为什么不是一段脚本"),
        *[rect(0.75 + i * 4.1, 1.9, 3.8, 3.35) for i in range(3)],
        *[text(1.0 + i * 4.1, 2.15, 3.3, 0.4, t, 18, bold=True, color=ACCENT)
          for i, t in enumerate(["结构分析", "因果化 / 撕裂", "一致初值 / 事件"])],
        *[text(1.0 + i * 4.1, 2.75, 3.3, 2.2, t, 14, color=INK, spacing=1.35)
          for i, t in enumerate([
              "BLT 分块、指标约简：\n先知道哪些方程要一起解。",
              "能显式算的显式算，\n剩下尽量小的隐式块。",
              "DAE 初值一致化；\n事件定位不靠猜。",
          ])],
        text(0.75, 5.55, 11.8, 0.5, "这些不是库函数，是编译链的一等公民。", 17, bold=True, color=DARK),
        *footer(5),
    ], "1 分 20 秒。这一页回答“为什么不用 scipy 写脚本”。重点是工程判断：结构、因果、事件决定成败。"))

    S.append(slide(WHITE, [
        *head("03 / 实测", "IEEE 300 潮流：4,604 个未知量"),
        *[
            text(0.75 + i * 4.05, 1.8, 3.6, 0.65, v, 34, bold=True, color=ACCENT, font="mono")
            for i, v in enumerate(["4,604", "0.32 s", "0 error"])
        ],
        *[
            text(0.75 + i * 4.05, 2.45, 3.6, 0.3, t, 11, color=GRAY)
            for i, t in enumerate(["展平未知量", "整进程墙钟（本机 release）", "XSD 校验"])
        ],
        code(0.75, 3.1, 8.3, 1.25, [
            '$ esml solve .../cases/case300.comp.esml -L library --json',
            '{"apiVersion":"esml.solve/v1",',
            ' "system":"Grid","unknowns":4604, ...}',
        ], 10.0),
        text(9.25, 3.2, 3.4, 1.9, "b1.net.Vm  1.0284201463\nb1.net.Va  0.1041501799\nb2.net.Vm  1.0353403863", 10.5, color=INK, font="mono", spacing=1.6),
        text(0.75, 4.5, 11.8, 0.5, "学术版 ≤16,384 未知量；IEEE 2869 / 1354 / 9241 需商业版。", 13, color=GRAY),
        text(0.75, 5.05, 11.8, 0.4, "学术版启动提示走 stderr；stdout 只有 JSON。", 12, color=GRAY),
        *footer(6),
    ], "2 分 30 秒。先给三个数字，再展示命令与输出。强调 0.32 秒是整进程墙钟，不是内核时间；学术版覆盖范围要主动说清楚。"))

    S.append(slide(WHITE, [
        *head("03 / 实测", "Python / Julia：一次性调用，MIT"),
        text(0.75, 1.72, 5.6, 0.3, "Python", 14, bold=True, color=ACCENT),
        text(6.9, 1.72, 5.6, 0.3, "Julia", 14, bold=True, color=ACCENT),
        code(0.75, 2.05, 5.8, 2.2, [
            'from esml import solve',
            '',
            'out = solve(',
            '    "case300.comp.esml",',
            '    libraries=["library"],',
            ')',
            'print(out.unknowns,',
            '      out.values["b1.net.Vm"])',
        ], 10.5),
        code(6.9, 2.05, 5.8, 2.2, [
            'using ESML',
            '',
            'out = solve(',
            '    "case300.comp.esml";',
            '    libraries=["library"],',
            ')',
            'println(out.unknowns,',
            '        out.values["b1.net.Vm"])',
        ], 10.5),
        text(0.75, 4.55, 11.8, 0.4, "只通过公开 JSON 协议调用引擎，不含引擎代码。", 16, bold=True, color=DARK),
        text(0.75, 5.05, 11.8, 0.3, 'pip install git+…@v0.1.0   ·   Pkg.add(url=…, rev="v0.1.0")', 10, color=GRAY, font="mono"),
        text(0.75, 5.55, 11.8, 0.3, "无状态：每次调用一个进程。", 12, color=GRAY),
        *footer(7),
    ], "1 分 20 秒。两边代码几乎一样：一次性调用、不保留状态。强调 MIT 与“不含引擎代码”，这是采用的关键。"))

    S.append(slide(WHITE, [
        *head("04 / MMPROPS", "物性：从状态点到相平衡"),
        *[rect(0.75 + i * 2.4, 1.9, 2.15, 2.75) for i in range(5)],
        *[text(0.92 + i * 2.4, 2.08, 1.9, 0.35, t, 18, bold=True, color=ACCENT, font="mono")
          for i, t in enumerate(["props", "sat", "flash", "vle", "envelope"])],
        *[text(0.92 + i * 2.4, 2.55, 1.9, 1.8, t, 11.5, color=INK, spacing=1.4)
          for i, t in enumerate([
              "T,p 状态点\nrho / h / s / cp",
              "饱和态\npsat / rho / hfg",
              "闪蒸\np,h 或 p,s",
              "二元 VLE\n泡点 / 露点",
              "T-x-y 包络\n固定 T",
          ])],
        text(0.75, 5.05, 11.8, 0.4, "数值可带单位；输出 JSON / CSV / 文本。", 16, bold=True, color=DARK),
        text(0.75, 5.55, 11.8, 0.4, "水/蒸汽、制冷剂、二元混合物；数据包可替换。", 13, color=GRAY),
        *footer(8),
    ], "1 分 40 秒。五个子命令一次讲完，不要展开；重点说“单位可带、输出可机读、失败响亮”。"))

    S.append(slide(WHITE, [
        *head("04 / MMPROPS", "水蒸汽与甲烷-乙烷：实测值"),
        table(0.75, 1.85, 7.7, 3.9, [
            ["量", "值"],
            ["水 rho（300 K / 1 bar）", "996.556338 kg/m3"],
            ["水 h", "112653.679689 J/kg"],
            ["水 psat（373.15 K）", "101417.996368 Pa"],
            ["闪蒸 T / x", "372.755929 K / 0.258033"],
            ["VLE p / y1（220 K）", "3993218.99 Pa / 0.847911"],
        ], [3.4, 4.3], 12),
        rect(8.7, 1.9, 3.9, 3.7, fill=MINT),
        text(8.95, 2.12, 3.4, 0.35, "设计取舍", 16, bold=True, color=ACCENT),
        text(8.95, 2.6, 3.4, 2.7, "两相无定义量返回 null，不是 0\n未收敛 / 超范围响亮失败\n坏数据包不退回内嵌基线\nVLE 只支持二元", 12.5, color=INK, spacing=1.35),
        text(0.75, 5.95, 11.8, 0.3, "所有数字可复跑：tutorials/water-steam。", 11, color=GRAY),
        *footer(9),
    ], "2 分 20 秒。挑两三个数字讲，不要念表；重点讲右栏的失败语义——这是工程用户最在意的。"))

    S.append(slide(WHITE, [
        *head("05 / 工程", "不是 demo：可复跑的工程判据"),
        *[
            text(0.75 + i * 3.0, 1.95, 2.8, 0.7, v, 32, bold=True, color=ACCENT, font="mono")
            for i, v in enumerate(["2,890", "23/23", "1,058", "29"])
        ],
        *[
            text(0.75 + i * 3.0, 2.65, 2.8, 0.6, t, 11, color=GRAY, spacing=1.3)
            for i, t in enumerate(["测试通过 / 0 失败", "合入闸门 G0–G22", "仓库内 .esml 模型", "crate 数"])
        ],
        text(0.75, 4.0, 11.8, 0.5, "每个数字都有日志和复跑命令。", 18, bold=True, color=DARK),
        text(0.75, 4.6, 11.8, 1.2, "全量测试 2890 passed / 0 failed / 32 ignored（271 个测试目标）；\n合入闸门 23/23；conformance 语料与 XSD 对拍。", 13, color=GRAY, spacing=1.4),
        *footer(10),
    ], "1 分 30 秒。这页建立信任：不是 demo，是可复跑的工程判据。数字要能报出来源。"))

    S.append(slide(WHITE, [
        *head("05 / 工程", "XML 模型，一条命令校验"),
        code(0.75, 1.9, 8.1, 1.75, [
            'for f in common compute basic; do',
            '  curl -O "https://docs.ai4energy.cn/schema/esml-canonical-$f.xsd"',
            'done',
            'xmllint --noout --schema \\',
            '  esml-canonical-basic.xsd model.esml',
        ], 10.0),
        text(9.2, 2.0, 3.4, 2.4, "• canonical 命名空间\n• MIT\n• IEEE 300：0 error", 13, color=INK, spacing=1.5),
        text(0.75, 4.05, 11.8, 0.5, "https://docs.ai4energy.cn/schema/", 20, bold=True, color=ACCENT, font="mono"),
        text(0.75, 4.8, 11.8, 0.8, "XML 工具、CI、模型仓库都能在入库前校验；\n不必启动求解器。", 14, color=GRAY, spacing=1.35),
        text(0.75, 5.7, 11.8, 0.3, "规范正文与语料：github.com/ai4energy/esml-language", 11, color=GRAY),
        *footer(11),
    ], "1 分钟。XSD 是标准工作的抓手：编辑器、CI、模型仓库都能用。强调 MIT 与实测 0 error。"))

    S.append(slide(WHITE, [
        *head("06 / 分发", "学术免费，商业授权；客户端 MIT"),
        rect(0.75, 1.9, 5.7, 3.0),
        text(1.0, 2.1, 5.2, 0.4, "学术版", 20, bold=True, color=ACCENT),
        text(1.0, 2.6, 5.2, 0.4, "≤ 16,384 展平未知量", 18, bold=True, color=DARK),
        text(1.0, 3.1, 5.2, 1.5, "引擎核心：solve / simulate /\nimport / check\n学术、教学、个人非商业免费", 12.5, color=INK, spacing=1.35),
        rect(6.9, 1.9, 5.7, 3.0, fill=LIGHT),
        text(7.15, 2.1, 5.2, 0.4, "商业版", 20, bold=True, color=ACCENT),
        text(7.15, 2.6, 5.2, 0.4, "无规模限制 + DTIL", 18, bold=True, color=DARK),
        text(7.15, 3.1, 5.2, 1.5, "常驻会话 / predict / snapshot\n商业用途需授权\nmingtao.li@gmail.com", 12.5, color=INK, spacing=1.35),
        text(0.75, 5.3, 11.8, 0.4, "下载：ai4energy/esml · ai4energy/mmprops", 14, bold=True, color=DARK),
        text(0.75, 5.8, 11.8, 0.6, "客户端：esml-python · ESML.jl · mmprops-python · MMProps.jl（均 MIT）", 11, color=GRAY),
        *footer(12),
    ], "1 分 10 秒。把边界讲清楚：学术版免费但有限；客户端 MIT；商业用途找邮箱。不要含糊。"))

    S.append(slide(WHITE, [
        *head("06 / 分发", "谁适合现在试"),
        *[rect(0.75 + i * 4.1, 1.9, 3.8, 2.8) for i in range(3)],
        *[text(1.0 + i * 4.1, 2.12, 3.3, 0.35, t, 18, bold=True, color=ACCENT)
          for i, t in enumerate(["课程 / 学生", "电力 / 能源", "过程 / 热力"])],
        *[text(1.0 + i * 4.1, 2.65, 3.3, 1.8, t, 13, color=INK, spacing=1.35)
          for i, t in enumerate([
              "IEEE 300、水蒸汽 / VLE、\n十分钟教程",
              "潮流、稳态 / 动态、\nXSD 入库校验",
              "物性、闪蒸、相包络、\n数据包",
          ])],
        text(0.75, 5.1, 11.8, 0.4, "第一英里问题按 bug 处理，24 小时内首响。", 16, bold=True, color=DARK),
        *footer(13),
    ], "1 分钟。给三类人各自一个入口；强调反馈承诺，让听众知道试了不会石沉大海。"))

    S.append(slide(WHITE, [
        *head("06 / 分发", "下一步：跑一个你自己的模型"),
        image(1.35, 2.0, 1.55, 1.55, HERE / "qr-esml.png"),
        image(5.35, 2.0, 1.55, 1.55, HERE / "qr-mmprops.png"),
        image(9.35, 2.0, 1.55, 1.55, HERE / "qr-feedback.png"),
        text(1.1, 3.7, 2.1, 0.3, "ESML 下载", 12, bold=True, color=DARK, align="c"),
        text(5.1, 3.7, 2.1, 0.3, "MMPROPS 下载", 12, bold=True, color=DARK, align="c"),
        text(9.1, 3.7, 2.1, 0.3, "反馈与访谈", 12, bold=True, color=DARK, align="c"),
        text(0.75, 4.5, 11.8, 0.45, "下载 → 教程 → 反馈", 22, bold=True, color=DARK),
        text(0.75, 5.1, 11.8, 0.8, "mingtao.li@gmail.com\n从分子到电网。", 14, color=GRAY, spacing=1.4),
        *footer(14),
    ], "1 分钟收尾。让听众扫一个码就好：ESML、MMPROPS、反馈。重复邮箱与 slogan，结束。"))

    S.append(slide(WHITE, [
        *head("备份", "命令清单"),
        code(0.75, 1.8, 11.8, 3.2, [
            '# ESML',
            'tar -xzf esml-academic-1.8.0-linux-x86_64.tar.gz',
            'export ESML_HOME=$PWD/esml-academic-1.8.0-linux-x86_64',
            '$ESML_HOME/esml solve \\',
            '  $ESML_HOME/library/domains/power/cases/case300.comp.esml \\',
            '  -L $ESML_HOME/library --json',
            '',
            '# MMPROPS',
            'tar -xzf mmprops-academic-0.6.0-linux-x86_64.tar.gz',
            'export MMPROPS_HOME=$PWD/mmprops-academic-0.6.0-linux-x86_64',
            '$MMPROPS_HOME/mmprops props --fluid water \\',
            '  --t "300 K" --p "1 bar" --format json',
        ], 10.0),
        text(0.75, 5.35, 11.8, 0.3, "教程：github.com/ai4energy/esml/tree/main/tutorials · 反馈：ai4energy.cn/feedback", 10.5, color=GRAY),
        *footer(15),
    ], "备份页，不占 20 分钟。Q&A 时直接投这页给命令。"))

    return S


# ── PPTX renderer ──────────────────────────────────────────────────────────
def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def set_run(run, size, bold, color, font):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = CJK if font == "cjk" else MONO
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", CJK if font == "cjk" else MONO)


def add_text(slide, el):
    box = slide.shapes.add_textbox(Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    align = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[el["align"]]
    lines = el["text"].split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = el.get("spacing", 1.15)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line
        set_run(r, el["size"], el["bold"], el["color"], el["font"])


def add_rect(slide, el):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if el["radius"] > 0 else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
    if el["radius"] > 0:
        try:
            shp.adjustments[0] = el["radius"]
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(el["fill"])
    if el["line"]:
        shp.line.color.rgb = rgb(el["line"])
        shp.line.width = Pt(el["line_w"])
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False


def add_code(slide, el):
    add_rect(slide, {"x": el["x"], "y": el["y"], "w": el["w"], "h": el["h"],
                     "fill": DARK2, "radius": 0.06, "line": None, "line_w": 0})
    box = slide.shapes.add_textbox(Inches(el["x"] + 0.18), Inches(el["y"] + 0.14),
                                   Inches(el["w"] - 0.3), Inches(el["h"] - 0.22))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, line in enumerate(el["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = line if line else " "
        set_run(r, el["size"], False, "D7EFE2", "mono")


def add_image(slide, el):
    slide.shapes.add_picture(el["path"], Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))


def add_line(slide, el):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(el["x1"]), Inches(el["y1"]),
                                      Inches(el["x2"]), Inches(el["y2"]))
    conn.line.color.rgb = rgb(el["color"])
    conn.line.width = Pt(el["width"])


def add_table(slide, el):
    rows = len(el["rows"]); cols = len(el["rows"][0])
    shape = slide.shapes.add_table(rows, cols, Inches(el["x"]), Inches(el["y"]), Inches(el["w"]), Inches(el["h"]))
    tbl = shape.table
    for i, w in enumerate(el["widths"]):
        tbl.columns[i].width = Inches(w)
    for r, row in enumerate(el["rows"]):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(MINT if r == 0 else (WHITE if r % 2 else LIGHT))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(value)
            set_run(run, el["size"], r == 0, DARK if r == 0 else INK, "cjk")


def build_pptx(slides):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(s["bg"])
        for el in s["elements"]:
            kind = el["t"]
            if kind == "text":
                add_text(sl, el)
            elif kind == "rect":
                add_rect(sl, el)
            elif kind == "code":
                add_code(sl, el)
            elif kind == "image":
                add_image(sl, el)
            elif kind == "line":
                add_line(sl, el)
            elif kind == "table":
                add_table(sl, el)
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(OUT_PPTX)


# ── PNG preview renderer ───────────────────────────────────────────────────
_font_cache = {}


def pil_font(size, bold=False, mono=False):
    key = (size, bold, mono)
    if key not in _font_cache:
        px = max(8, round(size * SCALE / 72))
        if mono:
            _font_cache[key] = ImageFont.truetype(MONO_TTF, px)
        else:
            path = CJK_BOLD if bold else CJK_REG
            try:
                _font_cache[key] = ImageFont.truetype(path, px, index=2)
            except Exception:
                _font_cache[key] = ImageFont.truetype(path, px)
    return _font_cache[key]


def wrap(draw, text, font, max_w):
    out = []
    for para in text.split("\n"):
        if not para:
            out.append("")
            continue
        line = ""
        for ch in para:
            if draw.textlength(line + ch, font=font) <= max_w:
                line += ch
            else:
                out.append(line)
                line = ch
        out.append(line)
    return out


def preview(slides):
    PREVIEW.mkdir(parents=True, exist_ok=True)
    thumbs = []
    for idx, s in enumerate(slides, 1):
        img = Image.new("RGB", (1600, 900), "#" + s["bg"])
        d = ImageDraw.Draw(img)
        for el in s["elements"]:
            kind = el["t"]
            if kind == "rect":
                xy = [el["x"] * SCALE, el["y"] * SCALE, (el["x"] + el["w"]) * SCALE, (el["y"] + el["h"]) * SCALE]
                r = int(el["radius"] * SCALE)
                if r > 0:
                    d.rounded_rectangle(xy, radius=r, fill="#" + el["fill"],
                                        outline=("#" + el["line"]) if el["line"] else None)
                else:
                    d.rectangle(xy, fill="#" + el["fill"],
                                outline=("#" + el["line"]) if el["line"] else None)
            elif kind == "line":
                d.line([el["x1"] * SCALE, el["y1"] * SCALE, el["x2"] * SCALE, el["y2"] * SCALE],
                       fill="#" + el["color"], width=max(1, round(el["width"] * SCALE / 72)))
            elif kind == "text":
                font = pil_font(el["size"], el["bold"], el["font"] == "mono")
                max_w = el["w"] * SCALE
                lines = wrap(d, el["text"], font, max_w)
                lh = el["size"] * SCALE / 72 * el.get("spacing", 1.15)
                y = el["y"] * SCALE
                for line in lines:
                    x = el["x"] * SCALE
                    if el["align"] == "c":
                        x = x + (max_w - d.textlength(line, font=font)) / 2
                    elif el["align"] == "r":
                        x = x + max_w - d.textlength(line, font=font)
                    d.text((x, y), line, font=font, fill="#" + el["color"])
                    y += lh
            elif kind == "code":
                xy = [el["x"] * SCALE, el["y"] * SCALE, (el["x"] + el["w"]) * SCALE, (el["y"] + el["h"]) * SCALE]
                d.rounded_rectangle(xy, radius=int(0.06 * SCALE), fill="#" + DARK2)
                font = pil_font(el["size"], False, True)
                y = (el["y"] + 0.14) * SCALE
                lh = el["size"] * SCALE / 72 * 1.12
                for line in el["lines"]:
                    d.text(((el["x"] + 0.18) * SCALE, y), line, font=font, fill="#D7EFE2")
                    y += lh
            elif kind == "table":
                x0, y0 = el["x"] * SCALE, el["y"] * SCALE
                widths = [w * SCALE for w in el["widths"]]
                row_h = el["h"] * SCALE / len(el["rows"])
                for r, row in enumerate(el["rows"]):
                    x = x0
                    for c, value in enumerate(row):
                        fill = "#EAF6F0" if r == 0 else ("#FFFFFF" if r % 2 else "#F4F7F5")
                        d.rectangle([x, y0 + r * row_h, x + widths[c], y0 + (r + 1) * row_h],
                                    fill=fill, outline="#DCE6E0")
                        font = pil_font(el["size"], r == 0)
                        d.text((x + 0.08 * SCALE, y0 + r * row_h + 0.05 * SCALE), str(value),
                               font=font, fill="#" + (DARK if r == 0 else INK))
                        x += widths[c]
            elif kind == "image":
                im = Image.open(el["path"]).convert("RGB")
                im = im.resize((round(el["w"] * SCALE), round(el["h"] * SCALE)), Image.LANCZOS)
                img.paste(im, (round(el["x"] * SCALE), round(el["y"] * SCALE)))
        path = PREVIEW / f"slide-{idx:02d}.png"
        img.save(path)
        thumbs.append(img.resize((400, 225), Image.LANCZOS))
    # contact sheet 4 × 4
    sheet = Image.new("RGB", (4 * 400, 4 * 225), "#FFFFFF")
    for i, th in enumerate(thumbs):
        sheet.paste(th, ((i % 4) * 400, (i // 4) * 225))
    sheet.save(PREVIEW / "contact-sheet.png")


if __name__ == "__main__":
    slides = build_slides()
    build_pptx(slides)
    preview(slides)
    print(f"wrote {OUT_PPTX} ({OUT_PPTX.stat().st_size / 1024:.1f} KiB)")
    print(f"preview: {PREVIEW}")
